import tkinter as tk
from tkinter import ttk, messagebox
import openpyxl
import os
import platform

# =============================================================================
# 1. KONFIGURATION & FARBEN
# =============================================================================
START_X = 1700
START_Y = 3000

# Spalten, die im rechten Info-Panel ausgeblendet werden sollen
HIDDEN_COLUMNS = ["X", "Y", "Pos", "TextPosition"]

THEMES = {
    "day": {
        "bg": "white",
        "grid": "#e0e0e0",
        "line": "#444444",
        "node_fill": "#0078D7", 
        "node_outline": "white",
        "text": "black",
        "highlight": "#FF8C00",   
        "hover": "#FF4040",       
        "panel_bg": "#f0f0f0",
        "panel_fg": "black"
    },
    "night": {
        "bg": "#1e1e1e", 
        "grid": "#333333",
        "line": "#cccccc", 
        "node_fill": "#4da6ff", 
        "node_outline": "#1e1e1e",
        "text": "#eeeeee",
        "highlight": "#ffaa00",   
        "hover": "#ff8888",       
        "panel_bg": "#2d2d30",
        "panel_fg": "white"
    }
}

# =============================================================================
# 2. DATEN LOGIK
# =============================================================================
class ExcelNetworkMap:
    def __init__(self, excel_path):
        self.excel_path = excel_path
        self.nodes = {} 
        self.edges = [] 
        self.max_x = 2000 
        self.max_y = 1500
        self.load_excel()

    def load_excel(self):
        self.nodes = {}
        self.edges = []
        
        if not os.path.exists(self.excel_path):
            self.create_demo_data()
            return True

        try:
            wb = openpyxl.load_workbook(self.excel_path, data_only=True)
            
            # --- KNOTEN ---
            if "Karte" in wb.sheetnames:
                ws = wb["Karte"]
                headers = {str(c.value).strip(): i for i, c in enumerate(ws[1]) if c.value}
                
                if all(k in headers for k in ["Kürzel", "X", "Y"]):
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        nid = row[headers["Kürzel"]]
                        x = row[headers["X"]]
                        y = row[headers["Y"]]
                        
                        if not nid or x is None or y is None: continue
                        
                        nid = str(nid).strip()
                        x, y = int(x), int(y)
                        
                        if x > self.max_x: self.max_x = x
                        if y > self.max_y: self.max_y = y
                        
                        name = row[headers["Name"]] if "Name" in headers else nid
                        typ = row[headers["Typ"]] if "Typ" in headers else ""
                        pos = str(row[headers["Pos"]]).lower() if "Pos" in headers else "s"

                        info_lines = [f"=== KNOTEN {nid} ==="]
                        for head, idx in headers.items():
                            if head in HIDDEN_COLUMNS: continue
                            val = row[idx]
                            if val is not None and str(val).strip() != "":
                                info_lines.append(f"• {head}: {val}")
                        
                        self.nodes[nid] = {
                            'label': f"{typ} {name}".strip() if typ else str(name),
                            'name': str(name).strip(),
                            'x': x, 'y': y, 'text_pos': pos, 
                            'info_text': "\n".join(info_lines)
                        }
            
            self.max_x = max(self.max_x, START_X + 1000)
            self.max_y = max(self.max_y, START_Y + 1000)

            # --- STRECKEN ---
            if "Strecken" in wb.sheetnames:
                ws = wb["Strecken"]
                h_map = {str(c.value).strip(): i for i, c in enumerate(ws[1]) if c.value}

                if "Start" in h_map and "Ende" in h_map:
                    edge_accum = {}
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        s_idx, e_idx = h_map["Start"], h_map["Ende"]
                        if row[s_idx] is None or row[e_idx] is None: continue
                        
                        s_id, e_id = str(row[s_idx]).strip(), str(row[e_idx]).strip()
                        key_parts = sorted([s_id, e_id])
                        edge_key = f"{key_parts[0]}_{key_parts[1]}"

                        track_details = []
                        raw_row_data = {}

                        for head, idx in h_map.items():
                            val = row[idx]
                            val_str = str(val) if val is not None else ""
                            raw_row_data[head] = val_str
                            
                            if val is not None and val_str.strip() != "":
                                if head in ["Start", "Ende"]: continue 
                                track_details.append(f"• {head}: {val}")
                        
                        info_block = "\n".join(track_details)
                        
                        if edge_key not in edge_accum:
                            edge_accum[edge_key] = {'start': s_id, 'end': e_id, 'blocks': [], 'raw_list': []}
                        
                        edge_accum[edge_key]['blocks'].append(info_block)
                        edge_accum[edge_key]['raw_list'].append(raw_row_data)

                    for k, d in edge_accum.items():
                        full_info = f"VERBINDUNG {d['start']} <-> {d['end']}\n" + ("="*30) + "\n"
                        for i, block in enumerate(d['blocks']):
                            full_info += f"\n--- Eintrag {i+1} ---\n{block}\n"
                        
                        self.edges.append({
                            'id': k, 
                            'start': d['start'], 
                            'end': d['end'], 
                            'info': full_info,
                            'raw_list': d['raw_list']
                        })
            wb.close()
            return True
            
        except PermissionError:
            messagebox.showerror("Fehler", "Die Excel-Datei ist geöffnet. Bitte schließen!")
            return False
        except Exception as e:
            print(f"Excel Error: {e}")
            self.create_demo_data()
            return False

    def create_demo_data(self):
        self.nodes = {
            "HAM": {"label": "Hamburg", "name": "Hamburg", "x": 500, "y": 200, "text_pos": "n", "info_text": "Demo"},
            "MUC": {"label": "München", "name": "München", "x": 1700, "y": 3000, "text_pos": "n", "info_text": "Demo"},
        }
        raw1 = {"VzG": "1000", "La-Seite1": "10", "La-Fplheft": "Heft A"}
        raw2 = {"VzG": "2000", "La-Seite1": "20", "La-Fplheft": "Heft B"}
        
        self.edges = [
            {"id": "1", "start": "HAM", "end": "MUC", "info": "Demo", "raw_list": [raw1, raw2]}
        ]
        self.max_x = 3500
        self.max_y = 3500

# =============================================================================
# 3. CANVAS LOGIC
# =============================================================================
class MapCanvas(tk.Canvas):
    def __init__(self, parent, **kwargs):
        super().__init__(parent, **kwargs)
        # FIX: Variable umbenannt, damit sie nicht den scale()-Befehl überschreibt
        self.zoom_level = 1.0  
        
        self.bind("<ButtonPress-1>", self.on_move_start)
        self.bind("<B1-Motion>", self.on_move_drag)
        self.bind("<MouseWheel>", self.on_mousewheel)
        self.bind("<Button-4>", lambda e: self.perform_zoom(e, 1.1))
        self.bind("<Button-5>", lambda e: self.perform_zoom(e, 0.9))
        self.bind("<Enter>", lambda e: self.focus_set())

    def on_move_start(self, event):
        self.scan_mark(event.x, event.y)

    def on_move_drag(self, event):
        self.scan_dragto(event.x, event.y, gain=1)

    def on_mousewheel(self, event):
        factor = 1.1 if event.delta > 0 else 0.9
        self.perform_zoom(event, factor)

    def perform_zoom(self, event, factor):
        x = self.canvasx(event.x)
        y = self.canvasy(event.y)
        
        # Berechnung mit der neuen Variable zoom_level
        new_scale = self.zoom_level * factor
        
        if new_scale < 0.1 or new_scale > 10: return
        
        self.zoom_level = new_scale
        
        # HIER war der Fehler: self.scale ist jetzt wieder die Original-Methode von Canvas
        self.scale("all", x, y, factor, factor)
        self.event_generate("<<ZoomChanged>>")
        
    def reset_zoom(self):
        self.zoom_level = 1.0

# =============================================================================
# 4. MAIN APP
# =============================================================================
class NetworkVisualizer:
    def __init__(self, root):
        self.root = root
        self.root.title("Netzplan Viewer - Pro")
        self.root.geometry("1400x900")
        
        self.excel_file = "configs/Netzdaten.xlsx"
        self.network = ExcelNetworkMap(self.excel_file)
        
        self.theme = THEMES["day"]
        self.current_theme_name = "day"
        self.selected_item = None
        self.hovered_item = None 
        
        self.setup_ui()
        self.draw_initial()
        self.root.after(100, lambda: self.center_view(START_X, START_Y))

    def setup_ui(self):
        toolbar = tk.Frame(self.root, bg="#f0f0f0", height=40)
        toolbar.pack(side=tk.TOP, fill=tk.X)
        
        tk.Button(toolbar, text="Tag/Nacht Modus", command=self.toggle_theme).pack(side=tk.LEFT, padx=10, pady=5)
        tk.Button(toolbar, text="🔄 Excel neu laden", command=self.reload_data, bg="#e1f5fe").pack(side=tk.LEFT, padx=10, pady=5)
        
        container = tk.PanedWindow(self.root, orient=tk.HORIZONTAL, sashwidth=5, bg="#d9d9d9")
        container.pack(fill=tk.BOTH, expand=True)
        
        self.map_frame = tk.Frame(container)
        container.add(self.map_frame, stretch="always")
        
        self.canvas = MapCanvas(self.map_frame, bg=self.theme["bg"], highlightthickness=0)
        self.canvas.pack(fill=tk.BOTH, expand=True)
        
        self.info_frame = tk.Frame(container, bg=self.theme["panel_bg"], width=350)
        container.add(self.info_frame, stretch="never")
        
        self.lbl_info_header = tk.Label(self.info_frame, text="Informationen", font=("Segoe UI", 12, "bold"),
                                        bg=self.theme["panel_bg"], fg=self.theme["panel_fg"])
        self.lbl_info_header.pack(pady=15)
        
        text_frame = tk.Frame(self.info_frame, bg=self.theme["panel_bg"])
        text_frame.pack(fill=tk.BOTH, expand=True, padx=2, pady=2)
        
        self.info_text = tk.Text(text_frame, width=35, font=("Consolas", 10), bd=0, padx=10, pady=10, wrap=tk.WORD)
        scrollbar = tk.Scrollbar(text_frame, command=self.info_text.yview)
        self.info_text.config(yscrollcommand=scrollbar.set)
        
        self.info_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.info_text.insert(tk.END, "Rechtsklick für Menü.\nLinksklick für Details.")
        self.info_text.config(state=tk.DISABLED)

        self.canvas.bind("<<ZoomChanged>>", self.on_zoom_changed)
        self.canvas.bind("<Button-1>", self.on_canvas_click, add="+")
        self.canvas.bind("<Motion>", self.on_mouse_move)
        
        self.canvas.bind("<Button-3>", self.on_right_click)
        if platform.system() == "Darwin":
            self.canvas.bind("<Button-2>", self.on_right_click)

    def reload_data(self):
        success = self.network.load_excel()
        if success:
            self.canvas.delete("all")
            self.canvas.reset_zoom()
            self.draw_initial()
            self.center_view(START_X, START_Y)
            self.show_toast("Daten erfolgreich neu geladen")

    def show_toast(self, message):
        toast = tk.Label(self.root, text=message, bg="#333333", fg="white", padx=20, pady=10, font=("Segoe UI", 10))
        toast.place(relx=0.5, rely=0.9, anchor="center")
        self.root.after(2000, toast.destroy)

    def copy_single_fahrplan_data(self, data, start_id, end_id):
        node_s = self.network.nodes.get(start_id, {})
        node_e = self.network.nodes.get(end_id, {})
        name_s = node_s.get('name', start_id)
        name_e = node_e.get('name', end_id)

        def get(key): return data.get(key, "").strip()

        text_block = ""
        
        page1 = get("La-Seite1")
        if page1:
            text_block += f"{name_s} - {name_e}\n"
            text_block += "Fahren Sie gemäß folgender Fahrplanangaben\n"
            text_block += f"Buchfahrplanheft {get('La-Fplheft')}, Seite {page1}, Spalte 2a/c, mit 30/50 km/h\n"

        page2 = get("La-Seite2")
        if page2:
            if text_block: text_block += "\n"
            text_block += f"{name_e} - {name_s}\n"
            text_block += "Fahren Sie gemäß folgender Fahrplanangaben\n"
            text_block += f"Buchfahrplanheft {get('La-Fplheft')}, Seite {page2}, Spalte 2a/c, mit 30/50 km/h\n"
        
        if not text_block:
            text_block = f"Keine Fahrplandaten (La-Seite1/2) für {name_s}-{name_e} gefunden."

        self.copy_to_clipboard(text_block, show_msg=False)
        self.show_toast(f"Daten für {get('VzG')} kopiert!")

    def on_right_click(self, event):
        tag, item_type = self.get_item_at_mouse(event.x, event.y)
        menu = tk.Menu(self.root, tearoff=0)
        
        if tag and item_type == "edge":
            edge_id = tag.split("_", 1)[1]
            edge = next((e for e in self.network.edges if e['id'] == edge_id), None)
            
            if edge:
                raw_list = edge.get('raw_list', [])
                for i, data in enumerate(raw_list):
                    vzg = data.get("VzG", f"Eintrag {i+1}")
                    if not vzg: vzg = f"Eintrag {i+1}"
                    
                    submenu = tk.Menu(menu, tearoff=0)
                    menu.add_cascade(label=f"Strecke {vzg}", menu=submenu)
                    
                    submenu.add_command(
                        label="📋 Fahrplandaten kopieren", 
                        command=lambda d=data: self.copy_single_fahrplan_data(d, edge['start'], edge['end'])
                    )
                menu.add_separator()
        
        # Koordinaten Berechnung angepasst auf self.canvas.zoom_level
        cx = self.canvas.canvasx(event.x)
        cy = self.canvas.canvasy(event.y)
        logic_x = int(cx / self.canvas.zoom_level) # FIX HIER
        logic_y = int(cy / self.canvas.zoom_level) # FIX HIER
        coord_str = f"{logic_x}, {logic_y}"
        
        menu.add_command(label=f"📍 Koord. kopieren: {coord_str}", 
                         command=lambda: [self.copy_to_clipboard(coord_str, show_msg=False), self.show_toast(f"Koordinaten kopiert: {coord_str}")])
        
        menu.add_separator()
        menu.add_command(label="Abbrechen")
        menu.post(event.x_root, event.y_root)

    def copy_to_clipboard(self, text, show_msg=True):
        self.root.clipboard_clear()
        self.root.clipboard_append(text)
        self.root.update()
        if show_msg:
            self.show_info(f"Kopiert:\n{text}")

    def center_view(self, target_x, target_y):
        self.canvas.update_idletasks()
        win_w = self.canvas.winfo_width()
        win_h = self.canvas.winfo_height()
        bbox = self.canvas.bbox("all")
        if not bbox: return
        x1, y1, x2, y2 = bbox
        total_w = x2 - x1
        total_h = y2 - y1
        if total_w == 0 or total_h == 0: return
        desired_left = target_x - (win_w / 2)
        desired_top = target_y - (win_h / 2)
        x_fraction = (desired_left - x1) / total_w
        y_fraction = (desired_top - y1) / total_h
        self.canvas.xview_moveto(x_fraction)
        self.canvas.yview_moveto(y_fraction)

    def toggle_theme(self):
        self.current_theme_name = "night" if self.current_theme_name == "day" else "day"
        self.theme = THEMES[self.current_theme_name]
        self.canvas.config(bg=self.theme["bg"])
        self.info_frame.config(bg=self.theme["panel_bg"])
        self.lbl_info_header.config(bg=self.theme["panel_bg"], fg=self.theme["panel_fg"])
        fg_text = "#eee" if self.current_theme_name == "night" else "black"
        bg_text = "#333" if self.current_theme_name == "night" else "white"
        self.info_text.config(bg=bg_text, fg=fg_text)
        self.info_text.master.config(bg=self.theme["panel_bg"]) 
        self.redraw_elements()

    def draw_initial(self):
        self.canvas.delete("all")
        t = self.theme
        nodes = self.network.nodes
        
        for edge in self.network.edges:
            s, e = edge['start'], edge['end']
            if s in nodes and e in nodes:
                x1, y1, x2, y2 = nodes[s]['x'], nodes[s]['y'], nodes[e]['x'], nodes[e]['y']
                tag = f"edge_{edge['id']}"
                self.canvas.create_line(x1, y1, x2, y2, width=2, fill=t["line"], 
                                      capstyle=tk.ROUND, tags=("elem", "edge", tag))

        r = 5
        for nid, n in nodes.items():
            cx, cy = n['x'], n['y']
            tag = f"node_{nid}"
            self.canvas.create_oval(cx-r, cy-r, cx+r, cy+r, 
                                    fill=t["node_fill"], outline=t["node_outline"], 
                                    tags=("elem", "node", tag))
            
            tx, ty, anch = cx, cy, "center"
            off = 12
            p = n.get('text_pos', 's')
            if p=='n': ty-=off; anch="s"
            elif p=='s': ty+=off; anch="n"
            elif p=='e': tx+=off; anch="w"
            elif p=='w': tx-=off; anch="e"
            self.canvas.create_text(tx, ty, text=n['label'], font=("Arial", 9), 
                                    fill=t["text"], tags=("elem", "label", tag), anchor=anch)

        self.canvas.config(scrollregion=self.canvas.bbox("all"))

    def get_item_at_mouse(self, screen_x, screen_y):
        cx = self.canvas.canvasx(screen_x)
        cy = self.canvas.canvasy(screen_y)
        screen_tolerance = 6 
        # FIX: scale -> zoom_level
        logic_tolerance = screen_tolerance / self.canvas.zoom_level
        items = self.canvas.find_overlapping(cx - logic_tolerance, cy - logic_tolerance, 
                                             cx + logic_tolerance, cy + logic_tolerance)
        prioritized = sorted(items, key=lambda i: 0 if "node" in self.canvas.gettags(i) else 1)
        for item in prioritized:
            tags = self.canvas.gettags(item)
            for tag in tags:
                if tag.startswith("node_") or tag.startswith("edge_"):
                    return tag, ("node" if tag.startswith("node_") else "edge")
        return None, None

    def update_item_style(self, tag, item_type, state):
        t = self.theme
        if state == "selected":
            col = t["highlight"]
            width = 4 if item_type == "edge" else 2
        elif state == "hover":
            col = t["hover"]
            width = 3 if item_type == "edge" else 2
        else:
            col = t["node_fill"] if item_type == "node" else t["line"]
            width = 2
        if item_type == "edge":
            self.canvas.itemconfig(tag, fill=col, width=width)
        elif item_type == "node":
            self.canvas.itemconfig(tag, fill=col)

    def on_mouse_move(self, event):
        tag, item_type = self.get_item_at_mouse(event.x, event.y)
        if tag == self.hovered_item: return
        if self.hovered_item and self.hovered_item != self.selected_item:
            old_type = "node" if "node_" in self.hovered_item else "edge"
            self.update_item_style(self.hovered_item, old_type, "normal")
        if tag and tag != self.selected_item:
            self.update_item_style(tag, item_type, "hover")
            self.canvas.config(cursor="hand2")
        else:
            self.canvas.config(cursor="")
        self.hovered_item = tag

    def on_canvas_click(self, event):
        tag, item_type = self.get_item_at_mouse(event.x, event.y)
        if tag:
            self.select_item(tag, item_type)
        else:
            if self.selected_item:
                 old_type = "node" if "node_" in self.selected_item else "edge"
                 self.update_item_style(self.selected_item, old_type, "normal")
            self.selected_item = None
            self.show_info("Rechtsklick für Menü.\nLinksklick für Details.")

    def select_item(self, tag, item_type):
        if self.selected_item and self.selected_item != tag:
            old_type = "node" if "node_" in self.selected_item else "edge"
            self.update_item_style(self.selected_item, old_type, "normal")
        self.selected_item = tag
        self.update_item_style(tag, item_type, "selected")
        info_text = ""
        uid = tag.split("_", 1)[1]
        if item_type == "node" and uid in self.network.nodes:
            info_text = self.network.nodes[uid]['info_text']
        elif item_type == "edge":
            for e in self.network.edges:
                if e['id'] == uid:
                    info_text = e['info']
                    break
        self.show_info(info_text)

    def on_zoom_changed(self, event):
        # FIX: scale -> zoom_level
        s = self.canvas.zoom_level
        state = tk.NORMAL if s > 0.5 else tk.HIDDEN
        self.canvas.itemconfig("label", state=state)

    def redraw_elements(self):
        t = self.theme
        self.canvas.itemconfig("edge", fill=t["line"], width=2)
        self.canvas.itemconfig("node", fill=t["node_fill"], outline=t["node_outline"])
        self.canvas.itemconfig("label", fill=t["text"])
        if self.selected_item:
            itype = "node" if "node_" in self.selected_item else "edge"
            self.update_item_style(self.selected_item, itype, "selected")

    def show_info(self, text):
        self.info_text.config(state=tk.NORMAL)
        self.info_text.delete(1.0, tk.END)
        self.info_text.insert(tk.END, text)
        self.info_text.config(state=tk.DISABLED)

if __name__ == "__main__":
    root = tk.Tk()
    app = NetworkVisualizer(root)
    root.mainloop()